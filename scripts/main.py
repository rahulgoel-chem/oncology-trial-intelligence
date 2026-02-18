import os
import glob
import requests
import pandas as pd
from datetime import datetime
from docx import Document
import smtplib
from email.mime.multipart import MIMEMultipart
from email.mime.base import MIMEBase
from email import encoders


# ============================================
# SETTINGS
# ============================================

URL = "https://clinicaltrials.gov/api/v2/studies"

PARAMS = {
    "query.cond": "Antibody Drug Conjugate OR ADC OR CAR-T OR lymphoma OR leukemia OR solid tumor",
    "pageSize": 1000
}

VERSION_FOLDER = "trial_versions"
CHANGE_FOLDER = "change_reports"
REPORT_FOLDER = "reports"

os.makedirs(VERSION_FOLDER, exist_ok=True)
os.makedirs(CHANGE_FOLDER, exist_ok=True)
os.makedirs(REPORT_FOLDER, exist_ok=True)

today_time = datetime.now().strftime("%Y-%m-%d_%H-%M")


# ============================================
# MODULE 1 — DOWNLOAD
# ============================================

all_studies = []
next_token = None

while True:

    if next_token:
        PARAMS["pageToken"] = next_token

    response = requests.get(URL, params=PARAMS)
    data = response.json()

    studies = data.get("studies", [])
    all_studies.extend(studies)

    next_token = data.get("nextPageToken")

    if not next_token:
        break

print("Downloaded:", len(all_studies))


# ============================================
# MODULE 2 — EXTRACT
# ============================================

trials = []

for study in all_studies:

    protocol = study.get("protocolSection", {})

    id_module = protocol.get("identificationModule", {})
    status_module = protocol.get("statusModule", {})
    sponsor_module = protocol.get("sponsorCollaboratorsModule", {})
    design_module = protocol.get("designModule", {})
    conditions_module = protocol.get("conditionsModule", {})

    arms_module = protocol.get("armsInterventionsModule", {})
    outcomes_module = protocol.get("outcomesModule", {})

    interventions = arms_module.get("interventions", [])
    intervention_names = [i.get("name") for i in interventions if i.get("name")]
    intervention_names = ", ".join(intervention_names)

    trials.append({
        "NCT Number": id_module.get("nctId"),
        "Study Title": id_module.get("briefTitle"),
        "Study Status": status_module.get("overallStatus"),
        "Sponsor": sponsor_module.get("leadSponsor", {}).get("name"),
        "Collaborators": str(sponsor_module.get("collaborators")),
        "Phases": str(design_module.get("phases")),
        "Enrollment": design_module.get("enrollmentInfo", {}).get("count"),
        "Funder Type": sponsor_module.get("leadSponsor", {}).get("class"),
        "Conditions": str(conditions_module.get("conditions")),
        "Interventions": intervention_names,
        "Start Date": status_module.get("startDateStruct", {}).get("date"),
        "Primary Completion Date": status_module.get("primaryCompletionDateStruct", {}).get("date"),
        "Completion Date": status_module.get("completionDateStruct", {}).get("date"),
        "Locations": str(protocol.get("contactsLocationsModule", {}))
    })

df = pd.DataFrame(trials)

file_path = f"{VERSION_FOLDER}/oncology_trials_{today_time}.csv"
df.to_csv(file_path, index=False)

print("Saved:", file_path)


# ============================================
# MODULE 3 — CHANGE DETECTION
# ============================================

files = sorted(glob.glob(f"{VERSION_FOLDER}/oncology_trials_*.csv"))

if len(files) < 2:
    print("Not enough files for comparison")
    exit()

OLD_FILE = files[-2]
NEW_FILE = files[-1]

old_df = pd.read_csv(OLD_FILE)
new_df = pd.read_csv(NEW_FILE)

old_df = old_df[old_df["Funder Type"] == "Industry"]
new_df = new_df[new_df["Funder Type"] == "Industry"]

KEY = "NCT Number"

old_df.set_index(KEY, inplace=True)
new_df.set_index(KEY, inplace=True)

columns_to_track = [
    "Study Status","Sponsor","Collaborators","Phases",
    "Enrollment","Start Date","Primary Completion Date",
    "Completion Date","Locations"
]

changes = []

common_ids = set(old_df.index).intersection(set(new_df.index))

for nct in common_ids:

    for col in columns_to_track:

        old_val = str(old_df.loc[nct].get(col))
        new_val = str(new_df.loc[nct].get(col))

        if old_val != new_val:

            changes.append({
                "NCT Number": nct,
                "Field": col,
                "Old": old_val,
                "New": new_val
            })

changes_df = pd.DataFrame(changes)

change_file = f"{CHANGE_FOLDER}/changes_{today_time}.csv"
changes_df.to_csv(change_file, index=False)


# ============================================
# MODULE 4B — DOC REPORT
# ============================================

doc = Document()
doc.add_heading("Oncology Trial Intelligence Update", 0)

for nct in changes_df["NCT Number"].unique():

    subset = changes_df[changes_df["NCT Number"] == nct]

    new_row = new_df.loc[nct]

    sponsor = new_row.get("Sponsor")
    intervention = new_row.get("Interventions")
    condition = new_row.get("Conditions")

    para = doc.add_paragraph()

    para.add_run(
        f"{sponsor}'s trial evaluating '{intervention}' "
        f"in patients with '{condition}' updated.\n"
    ).bold = True

    para.add_run(f"{nct} — {today_time}\n")

    for _, row in subset.iterrows():
        doc.add_paragraph(
            f"{row['Field']} changed from '{row['Old']}' to '{row['New']}'",
            style="List Bullet"
        )

report_path = f"{REPORT_FOLDER}/report_{today_time}.docx"
doc.save(report_path)

print("Report generated:", report_path)


# ============================================
# MODULE 5 — EMAIL
# ============================================

EMAIL_USER = os.getenv("EMAIL_USER")
EMAIL_PASS = os.getenv("EMAIL_PASS")
EMAIL_TO = os.getenv("EMAIL_TO")

if EMAIL_USER:

    msg = MIMEMultipart()
    msg["From"] = EMAIL_USER
    msg["To"] = EMAIL_TO
    msg["Subject"] = "Oncology Intelligence Report"

    part = MIMEBase("application", "octet-stream")
    with open(report_path, "rb") as f:
        part.set_payload(f.read())

    encoders.encode_base64(part)
    part.add_header(
        "Content-Disposition",
        f"attachment; filename={os.path.basename(report_path)}"
    )

    msg.attach(part)

    server = smtplib.SMTP_SSL("smtp.gmail.com", 465)
    server.login(EMAIL_USER, EMAIL_PASS)
    server.send_message(msg)
    server.quit()

    print("Email sent")
